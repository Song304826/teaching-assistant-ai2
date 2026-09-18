from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt


BLUE = PptRGBColor(36, 87, 214)
NAVY = PptRGBColor(19, 42, 86)
INK = PptRGBColor(24, 33, 51)
MUTED = PptRGBColor(91, 105, 130)
PALE = PptRGBColor(241, 245, 252)
WHITE = PptRGBColor(255, 255, 255)
CYAN = PptRGBColor(43, 198, 214)
VIOLET = PptRGBColor(124, 92, 246)
CORAL = PptRGBColor(255, 102, 102)
AMBER = PptRGBColor(255, 184, 77)
MINT = PptRGBColor(45, 196, 145)
SOFT_BLUE = PptRGBColor(229, 238, 255)
SOFT_VIOLET = PptRGBColor(239, 234, 255)
SOFT_CORAL = PptRGBColor(255, 235, 235)
SOFT_MINT = PptRGBColor(229, 248, 241)


def safe_filename(value: str, fallback: str = "教学方案") -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|\r\n]+", "_", str(value)).strip(" ._")
    return cleaned[:50] or fallback


def _set_run_font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_label_paragraph(document: Document, label: str, value: Any) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    label_run = paragraph.add_run(label)
    _set_run_font(label_run, "Microsoft YaHei", 10.5, True)
    value_run = paragraph.add_run(str(value or ""))
    _set_run_font(value_run, "SimSun", 10.5)


def create_lesson_docx(plan: dict[str, Any], requirements: dict[str, str]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(str(plan.get("lesson_title") or requirements.get("课题") or "教学设计"))
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("课研助手生成的第一版教学方案")
    _set_run_font(run, "Microsoft YaHei", 10)
    run.font.color.rgb = RGBColor(90, 100, 118)

    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.55)
    metadata = [
        ("学科年级", requirements.get("学科年级", "")),
        ("教材版本", requirements.get("教材版本", "")),
        ("课时时长", requirements.get("课时时长", "")),
        ("学生学情", requirements.get("学生学情", "")),
        ("互动与风格", requirements.get("互动与风格", "")),
    ]
    for label, value in metadata:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Inches(1.35), Inches(5.55)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(cells[0], "EAF0FB")
        for cell in cells:
            _set_cell_margins(cell)
        left = cells[0].paragraphs[0]
        right = cells[1].paragraphs[0]
        _set_run_font(left.add_run(label), "Microsoft YaHei", 10, True)
        _set_run_font(right.add_run(str(value)), "SimSun", 10)

    document.add_paragraph()
    document.add_heading("一 教学设计思路", level=1)
    document.add_paragraph(str(plan.get("summary", "")))

    document.add_heading("二 教学目标", level=1)
    for item in plan.get("teaching_goals", []):
        paragraph = document.add_paragraph(style="List Bullet")
        _set_run_font(paragraph.add_run(str(item)), "SimSun", 10.5)

    document.add_heading("三 教学重点与难点", level=1)
    for item in plan.get("key_points", []):
        _add_label_paragraph(document, "教学重点：", item)
    for item in plan.get("difficult_points", []):
        _add_label_paragraph(document, "教学难点：", item)

    document.add_heading("四 教学过程", level=1)
    for index, item in enumerate(plan.get("teaching_process", []), start=1):
        phase = item.get("phase", f"教学环节{index}")
        duration = item.get("duration", "")
        document.add_heading(f"{index} {phase}  {duration}", level=2)
        _add_label_paragraph(document, "教师活动：", item.get("teacher", ""))
        _add_label_paragraph(document, "学生活动：", item.get("student", ""))
        _add_label_paragraph(document, "设计意图：", item.get("purpose", ""))

    activity = plan.get("interactive_activity", {})
    document.add_heading("五 课堂互动", level=1)
    _add_label_paragraph(document, "活动名称：", activity.get("name", ""))
    _add_label_paragraph(document, "活动规则：", activity.get("rules", ""))
    _add_label_paragraph(document, "评价反馈：", activity.get("feedback", ""))

    homework = plan.get("homework", {})
    document.add_heading("六 分层作业", level=1)
    _add_label_paragraph(document, "基础任务：", homework.get("basic", ""))
    _add_label_paragraph(document, "提升任务：", homework.get("advanced", ""))

    references = plan.get("references", [])
    if references:
        document.add_heading("七 参考资料", level=1)
        for item in references:
            document.add_paragraph(str(item), style="List Bullet")

    document.core_properties.title = str(plan.get("lesson_title", "教学设计"))
    document.core_properties.subject = "课堂教学设计"
    document.core_properties.author = "课研助手"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _new_handout(title_text: str, subtitle_text: str) -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.line_spacing = 1.35
    for style_name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(26, 58, 120)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(title_text)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(subtitle_text)
    _set_run_font(run, "Microsoft YaHei", 9.5)
    run.font.color.rgb = RGBColor(95, 108, 133)
    return document


def _save_document(document: Document, subject: str) -> bytes:
    document.core_properties.title = subject
    document.core_properties.author = "课研助手"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def create_teacher_handout_docx(plan: dict[str, Any], requirements: dict[str, str]) -> bytes:
    lesson = str(plan.get("lesson_title") or requirements.get("课题") or "教学讲义")
    document = _new_handout(f"{lesson}｜教师教学讲义", "备课提示、例题讲解与课堂追问建议")
    handout = plan.get("teacher_handout", {}) or {}
    document.add_heading("一 课程导航", level=1)
    _add_label_paragraph(document, "教学顺序：", requirements.get("教学顺序", ""))
    _add_label_paragraph(document, "时间安排：", f"总课时 {requirements.get('课时时长', '')}；课件覆盖 {requirements.get('PPT覆盖时长', '')}")
    _add_label_paragraph(document, "讲解线索：", handout.get("overview") or plan.get("summary", ""))
    fusion = plan.get("interdisciplinary_design", {}) or {}
    if fusion.get("subject") and fusion.get("subject") != "不启用":
        document.add_heading("二 跨学科导入", level=1)
        _add_label_paragraph(document, "融合学科：", fusion.get("subject", ""))
        _add_label_paragraph(document, "连接点：", fusion.get("connection", ""))
        _add_label_paragraph(document, "课堂任务：", fusion.get("activity", ""))
    document.add_heading("三 核心知识与讲解提示", level=1)
    points = handout.get("knowledge_points") or plan.get("key_points", [])
    for item in points:
        document.add_paragraph(str(item), style="List Bullet")
    document.add_heading("四 例题与讲解", level=1)
    examples = handout.get("examples", [])
    if not examples:
        examples = [{"question": item.get("question", ""), "solution": item.get("explanation", ""), "teaching_note": "引导学生说明判断依据。"} for item in plan.get("quiz", [])]
    for index, item in enumerate(examples, start=1):
        document.add_heading(f"例题 {index}", level=2)
        _add_label_paragraph(document, "题目：", item.get("question", ""))
        _add_label_paragraph(document, "解答：", item.get("solution", ""))
        _add_label_paragraph(document, "讲解提示：", item.get("teaching_note", ""))
    document.add_heading("五 易错提醒", level=1)
    misconceptions = handout.get("misconceptions") or plan.get("difficult_points", [])
    for item in misconceptions:
        document.add_paragraph(str(item), style="List Bullet")
    return _save_document(document, f"{lesson}教师讲义")


def create_student_handout_docx(plan: dict[str, Any], requirements: dict[str, str]) -> bytes:
    lesson = str(plan.get("lesson_title") or requirements.get("课题") or "课堂学习单")
    document = _new_handout(f"{lesson}｜学生课堂学习单", "姓名：____________  班级：____________  日期：____________")
    handout = plan.get("student_handout", {}) or {}
    document.add_heading("一 学习目标", level=1)
    for item in handout.get("learning_goals") or plan.get("teaching_goals", []):
        document.add_paragraph("□ " + str(item))
    document.add_heading("二 课堂笔记", level=1)
    for item in handout.get("key_notes") or plan.get("key_points", []):
        document.add_paragraph("• " + str(item) + "\n  我的补充：____________________________________________")
    document.add_heading("三 跟我做例题", level=1)
    examples = handout.get("guided_examples", [])
    for index, item in enumerate(examples, start=1):
        document.add_heading(f"例题 {index}", level=2)
        document.add_paragraph(str(item.get("question", "")))
        document.add_paragraph("解题步骤提示：" + str(item.get("steps", "")))
        document.add_paragraph("我的过程：\n________________________________________________________\n________________________________________________________")
    document.add_heading("四 当堂练习", level=1)
    practice = handout.get("practice") or [item.get("question", "") for item in plan.get("quiz", [])]
    for index, item in enumerate(practice, start=1):
        document.add_paragraph(f"{index}. {item}\n   答：__________________________________________________")
    document.add_heading("五 学习反思", level=1)
    reflections = handout.get("reflection") or ["我能否用自己的话说出今天的方法？", "我最容易出错的步骤是什么？"]
    for item in reflections:
        document.add_paragraph("□ " + str(item))
    return _save_document(document, f"{lesson}学生学习单")


def _add_textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = PptInches(0.06)
    frame.margin_right = PptInches(0.06)
    frame.margin_top = PptInches(0.04)
    frame.margin_bottom = PptInches(0.04)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = "Microsoft YaHei"
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _set_slide_background(slide, color) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _content_lines(content: Any) -> list[str]:
    text = re.sub(r"\s+", " ", str(content or "")).strip()
    parts = [part.strip() for part in re.split(r"[；;。]\s*", text) if part.strip()]
    return parts[:5] or [text]


def _add_shape(slide, shape_type, left, top, width, height, fill_color, radius_line=None):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if radius_line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = radius_line
    return shape


def _add_pill(slide, left, top, width, text, fill_color, text_color=WHITE, size=11):
    _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, PptInches(0.38), fill_color)
    _add_textbox(slide, left, top, width, PptInches(0.38), text, size, text_color, True, PP_ALIGN.CENTER)


def _add_slide_footer(slide, presentation, index: int, total: int, title: str, dark: bool = False):
    base = PptRGBColor(52, 70, 105) if dark else PptRGBColor(210, 219, 235)
    active = CYAN if dark else BLUE
    _add_textbox(
        slide,
        PptInches(0.72),
        PptInches(7.02),
        PptInches(7.5),
        PptInches(0.24),
        title,
        8.5,
        PptRGBColor(174, 193, 226) if dark else MUTED,
    )
    _add_textbox(
        slide,
        PptInches(11.85),
        PptInches(6.98),
        PptInches(0.72),
        PptInches(0.28),
        f"{index:02d}/{total:02d}",
        9,
        PptRGBColor(174, 193, 226) if dark else MUTED,
        True,
        PP_ALIGN.RIGHT,
    )
    track_left = PptInches(8.55)
    track_width = PptInches(3.0)
    _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, track_left, PptInches(7.10), track_width, PptInches(0.06), base)
    progress = max(0.08, min(1.0, index / max(total, 1)))
    _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, track_left, PptInches(7.10), int(track_width * progress), PptInches(0.06), active)


def _add_bullet_list(slide, left, top, width, height, lines, color=INK, size=18, accent=BLUE):
    if not lines:
        return
    row_height = height / max(len(lines), 1)
    for line_index, line in enumerate(lines):
        y = top + row_height * line_index
        _add_shape(slide, MSO_SHAPE.OVAL, left, y + PptInches(0.15), PptInches(0.12), PptInches(0.12), accent)
        _add_textbox(slide, left + PptInches(0.28), y, width - PptInches(0.28), row_height, line, size, color)


def _add_decorative_orbs(slide, dark: bool = True):
    colors = (VIOLET, CYAN, BLUE) if dark else (SOFT_VIOLET, SOFT_BLUE, SOFT_MINT)
    _add_shape(slide, MSO_SHAPE.OVAL, PptInches(11.15), PptInches(-0.65), PptInches(2.6), PptInches(2.6), colors[0])
    _add_shape(slide, MSO_SHAPE.OVAL, PptInches(11.72), PptInches(0.15), PptInches(1.5), PptInches(1.5), colors[1])
    _add_shape(slide, MSO_SHAPE.OVAL, PptInches(-0.6), PptInches(6.35), PptInches(1.45), PptInches(1.45), colors[2])


def _clean_title(value: Any, fallback: str) -> str:
    title = re.sub(r"\s+", " ", str(value or fallback)).strip()
    return title[:32] if title else fallback


def _agenda_slide(presentation, blank, outlines, lesson_title):
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, PptRGBColor(248, 250, 255))
    _add_decorative_orbs(slide, False)
    _add_pill(slide, PptInches(0.78), PptInches(0.55), PptInches(1.45), "LEARNING MAP", BLUE, WHITE, 9)
    _add_textbox(slide, PptInches(0.78), PptInches(1.02), PptInches(8.2), PptInches(0.72), "今天我们这样学", 29, INK, True)
    _add_textbox(
        slide,
        PptInches(0.8),
        PptInches(1.76),
        PptInches(8.5),
        PptInches(0.42),
        f"围绕核心问题逐步推进，共 {len(outlines)} 个学习环节",
        14,
        MUTED,
    )
    palette = [(SOFT_BLUE, BLUE), (SOFT_VIOLET, VIOLET), (SOFT_MINT, MINT), (SOFT_CORAL, CORAL), (PptRGBColor(255, 244, 222), AMBER), (PptRGBColor(232, 246, 252), CYAN)]
    for idx, item in enumerate(outlines[:6]):
        row, col = divmod(idx, 3)
        x = PptInches(0.78 + col * 4.12)
        y = PptInches(2.45 + row * 1.72)
        bg, accent = palette[idx]
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, PptInches(3.75), PptInches(1.34), bg)
        _add_textbox(slide, x + PptInches(0.25), y + PptInches(0.20), PptInches(0.48), PptInches(0.42), f"{idx + 1:02d}", 17, accent, True)
        _add_textbox(
            slide,
            x + PptInches(0.82),
            y + PptInches(0.16),
            PptInches(2.62),
            PptInches(0.95),
            _clean_title(item.get("title"), f"学习环节{idx + 1}"),
            15,
            INK,
            True,
        )
    if len(outlines) > 6:
        _add_textbox(slide, PptInches(0.82), PptInches(6.05), PptInches(8.6), PptInches(0.35), f"后续还有 {len(outlines) - 6} 个巩固与应用环节", 11, MUTED)
    _add_slide_footer(slide, presentation, 2, len(outlines) + 3, lesson_title)


def _standard_slide_header(slide, index, title, tag, dark=False, accent=BLUE):
    title_color = WHITE if dark else INK
    muted_color = PptRGBColor(178, 197, 228) if dark else MUTED
    _add_textbox(slide, PptInches(0.78), PptInches(0.56), PptInches(0.75), PptInches(0.42), f"{index:02d}", 13, accent, True)
    _add_textbox(slide, PptInches(1.55), PptInches(0.54), PptInches(9.8), PptInches(0.82), title, 27, title_color, True)
    _add_textbox(slide, PptInches(1.57), PptInches(1.32), PptInches(5.6), PptInches(0.30), tag, 9, muted_color, True)


def _layout_focus(slide, lines, accent):
    first = lines[0]
    rest = lines[1:]
    _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(0.78), PptInches(1.92), PptInches(11.76), PptInches(1.62), accent)
    _add_textbox(slide, PptInches(1.10), PptInches(2.13), PptInches(11.05), PptInches(1.15), first, 24 if len(first) < 42 else 20, WHITE, True, PP_ALIGN.CENTER)
    if rest:
        _add_bullet_list(slide, PptInches(1.12), PptInches(3.90), PptInches(10.8), PptInches(2.45), rest, INK, 17, accent)


def _layout_cards(slide, lines, palette):
    for idx, line in enumerate(lines[:4]):
        row, col = divmod(idx, 2)
        x = PptInches(0.82 + col * 6.03)
        y = PptInches(1.95 + row * 2.18)
        bg, accent = palette[idx % len(palette)]
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, PptInches(5.62), PptInches(1.80), bg)
        _add_textbox(slide, x + PptInches(0.26), y + PptInches(0.22), PptInches(0.55), PptInches(0.42), f"{idx + 1:02d}", 15, accent, True)
        _add_textbox(slide, x + PptInches(0.92), y + PptInches(0.18), PptInches(4.35), PptInches(1.30), line, 16 if len(line) < 55 else 14, INK, True)


def _layout_process(slide, lines):
    items = lines[:4]
    width = 11.45 / max(len(items), 1)
    colors = [(BLUE, SOFT_BLUE), (VIOLET, SOFT_VIOLET), (CYAN, PptRGBColor(228, 248, 251)), (MINT, SOFT_MINT)]
    for idx, line in enumerate(items):
        accent, bg = colors[idx]
        x = PptInches(0.82 + idx * width)
        card_w = PptInches(width - 0.18)
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, PptInches(2.02), card_w, PptInches(3.86), bg)
        _add_shape(slide, MSO_SHAPE.OVAL, x + PptInches(0.28), PptInches(2.38), PptInches(0.72), PptInches(0.72), accent)
        _add_textbox(slide, x + PptInches(0.28), PptInches(2.38), PptInches(0.72), PptInches(0.72), str(idx + 1), 18, WHITE, True, PP_ALIGN.CENTER)
        _add_textbox(slide, x + PptInches(0.27), PptInches(3.30), card_w - PptInches(0.54), PptInches(2.12), line, 15 if len(line) < 46 else 13, INK, True)


def _layout_misconception(slide, lines):
    midpoint = max(1, (len(lines) + 1) // 2)
    groups = [("容易出错", lines[:midpoint], SOFT_CORAL, CORAL), ("正确理解", lines[midpoint:] or ["利用等式性质检查每一步"], SOFT_MINT, MINT)]
    for idx, (label, items, bg, accent) in enumerate(groups):
        x = PptInches(0.82 + idx * 6.05)
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, PptInches(1.92), PptInches(5.62), PptInches(4.48), bg)
        _add_pill(slide, x + PptInches(0.26), PptInches(2.18), PptInches(1.34), label, accent, WHITE, 11)
        _add_bullet_list(slide, x + PptInches(0.35), PptInches(2.92), PptInches(4.86), PptInches(2.96), items, INK, 16, accent)


def _layout_challenge(slide, lines):
    _set_slide_background(slide, NAVY)
    _add_decorative_orbs(slide, True)
    _add_pill(slide, PptInches(0.82), PptInches(1.88), PptInches(1.48), "CLASS CHALLENGE", VIOLET, WHITE, 9)
    _add_textbox(slide, PptInches(0.82), PptInches(2.38), PptInches(11.45), PptInches(1.38), lines[0], 28 if len(lines[0]) < 38 else 22, WHITE, True, PP_ALIGN.CENTER)
    if len(lines) > 1:
        for idx, line in enumerate(lines[1:4]):
            x = PptInches(1.02 + idx * 3.88)
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, PptInches(4.30), PptInches(3.48), PptInches(1.32), PptRGBColor(39, 66, 118))
            _add_textbox(slide, x + PptInches(0.18), PptInches(4.47), PptInches(3.12), PptInches(0.92), line, 14, WHITE, True, PP_ALIGN.CENTER)


def _closing_slide(presentation, blank, plan, requirements, total):
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, NAVY)
    _add_decorative_orbs(slide, True)
    _add_pill(slide, PptInches(0.82), PptInches(0.58), PptInches(1.55), "LESSON COMPLETE", CYAN, NAVY, 9)
    _add_textbox(slide, PptInches(0.82), PptInches(1.25), PptInches(8.8), PptInches(0.86), "把方法带走，把思考留下", 31, WHITE, True)
    activity = plan.get("interactive_activity", {}) or {}
    homework = plan.get("homework", {}) or {}
    blocks = [
        ("课堂互动", activity.get("name") or activity.get("rules") or "用自己的话解释今天的方法", VIOLET),
        ("课后巩固", homework.get("basic") or "完成基础练习并检查关键步骤", CYAN),
        ("进阶挑战", homework.get("advanced") or "尝试解决一个变式问题", AMBER),
    ]
    for idx, (label, value, accent) in enumerate(blocks):
        x = PptInches(0.82 + idx * 4.08)
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, PptInches(2.58), PptInches(3.72), PptInches(2.55), PptRGBColor(39, 66, 118))
        _add_textbox(slide, x + PptInches(0.25), PptInches(2.88), PptInches(3.20), PptInches(0.36), label, 12, accent, True)
        _add_textbox(slide, x + PptInches(0.25), PptInches(3.38), PptInches(3.20), PptInches(1.26), value, 16 if len(str(value)) < 42 else 13, WHITE, True)
    meta = f"{requirements.get('学科年级', '')}  ·  {requirements.get('课时时长', '')}"
    _add_textbox(slide, PptInches(0.85), PptInches(6.22), PptInches(5.5), PptInches(0.38), meta, 11, PptRGBColor(175, 195, 229))
    _add_slide_footer(slide, presentation, total, total, str(plan.get("lesson_title", "教学课件")), True)


PPT_STYLE_MAP = {
    "科技蓝": {
        "132A56": "132A56", "2457D6": "2457D6", "2BC6D6": "2BC6D6", "7C5CF6": "7C5CF6",
    },
    "活力暖色": {
        "132A56": "5A2434", "2457D6": "E85D3F", "2BC6D6": "FFB84D", "7C5CF6": "C64F76",
        "E5EEFF": "FFF0E8", "EFEAFF": "FCE8EE", "E5F8F1": "FFF5DC", "F8FAFF": "FFFBF7",
    },
    "清新自然": {
        "132A56": "164D45", "2457D6": "268D78", "2BC6D6": "54B9A8", "7C5CF6": "6B8E62",
        "E5EEFF": "E6F5F1", "EFEAFF": "EDF4E8", "E5F8F1": "E2F6ED", "F8FAFF": "F7FCF9",
    },
    "简约学术": {
        "132A56": "252B36", "2457D6": "3B5B92", "2BC6D6": "6F87A8", "7C5CF6": "596579",
        "E5EEFF": "E9EEF5", "EFEAFF": "EEF0F4", "E5F8F1": "EDF2F0", "F8FAFF": "FBFCFE",
    },
}


def _apply_ppt_style(presentation: Presentation, style_name: str) -> None:
    mapping = PPT_STYLE_MAP.get(style_name, PPT_STYLE_MAP["科技蓝"])

    def mapped(rgb):
        if rgb is None:
            return None
        code = str(rgb).upper()
        replacement = mapping.get(code)
        return PptRGBColor.from_string(replacement) if replacement else None

    for slide in presentation.slides:
        for shape in slide.shapes:
            try:
                replacement = mapped(shape.fill.fore_color.rgb)
                if replacement:
                    shape.fill.fore_color.rgb = replacement
            except (AttributeError, TypeError, ValueError):
                pass
            try:
                replacement = mapped(shape.line.color.rgb)
                if replacement:
                    shape.line.color.rgb = replacement
            except (AttributeError, TypeError, ValueError):
                pass
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    try:
                        replacement = mapped(run.font.color.rgb)
                        if replacement:
                            run.font.color.rgb = replacement
                    except (AttributeError, TypeError, ValueError):
                        pass


def create_lesson_pptx(plan: dict[str, Any], requirements: dict[str, str]) -> bytes:
    presentation = Presentation()
    presentation.slide_width = PptInches(13.333)
    presentation.slide_height = PptInches(7.5)
    blank = presentation.slide_layouts[6]

    title = str(plan.get("lesson_title") or requirements.get("课题") or "教学课件")
    outlines = plan.get("ppt_outline", []) or []
    total = len(outlines) + 3

    cover = presentation.slides.add_slide(blank)
    _set_slide_background(cover, NAVY)
    _add_decorative_orbs(cover, True)
    _add_pill(cover, PptInches(0.86), PptInches(0.68), PptInches(1.72), "AI LESSON DESIGN", CYAN, NAVY, 9)
    _add_textbox(cover, PptInches(0.86), PptInches(1.47), PptInches(10.4), PptInches(1.30), title, 37 if len(title) < 20 else 31, WHITE, True)
    _add_textbox(cover, PptInches(0.90), PptInches(2.92), PptInches(9.55), PptInches(1.12), plan.get("summary", ""), 17, PptRGBColor(214, 225, 246))
    meta_items = [requirements.get("学科年级", ""), requirements.get("教材版本", ""), requirements.get("课时时长", "")]
    x = 0.90
    for value in [item for item in meta_items if item]:
        width = min(2.45, max(1.25, 0.34 * len(str(value)) + 0.55))
        _add_pill(cover, PptInches(x), PptInches(5.32), PptInches(width), str(value), PptRGBColor(39, 66, 118), WHITE, 11)
        x += width + 0.18
    _add_textbox(cover, PptInches(0.90), PptInches(6.35), PptInches(7.0), PptInches(0.34), "课研助手 · 课堂教学方案", 10, PptRGBColor(174, 195, 232))

    _agenda_slide(presentation, blank, outlines, title)

    light_palette = [(SOFT_BLUE, BLUE), (SOFT_VIOLET, VIOLET), (SOFT_MINT, MINT), (SOFT_CORAL, CORAL)]
    for index, item in enumerate(outlines, start=1):
        slide = presentation.slides.add_slide(blank)
        page_title = _clean_title(item.get("title"), f"学习环节{index}")
        lines = _content_lines(item.get("content", ""))
        lowered = page_title + " " + " ".join(lines)
        is_challenge = any(word in lowered for word in ("练习", "训练", "挑战", "游戏", "互动", "板演"))
        is_misconception = any(word in lowered for word in ("易错", "错误", "辨一辨", "误区"))
        is_process = any(word in lowered for word in ("步骤", "法则", "方法", "流程", "怎样", "怎么"))

        if is_challenge:
            _layout_challenge(slide, lines)
            _standard_slide_header(slide, index, page_title, "THINK · DISCUSS · SHARE", True, CYAN)
            _add_slide_footer(slide, presentation, index + 2, total, title, True)
        else:
            _set_slide_background(slide, PptRGBColor(248, 250, 255))
            _add_decorative_orbs(slide, False)
            accent = (BLUE, VIOLET, CYAN, MINT)[(index - 1) % 4]
            tag = "KEY CONCEPT" if index % 3 == 1 else ("METHOD" if index % 3 == 2 else "CHECKPOINT")
            _standard_slide_header(slide, index, page_title, tag, False, accent)
            if is_misconception:
                _layout_misconception(slide, lines)
            elif is_process or index % 3 == 0:
                _layout_process(slide, lines)
            elif index % 3 == 1:
                _layout_focus(slide, lines, accent)
            else:
                _layout_cards(slide, lines, light_palette)
            _add_slide_footer(slide, presentation, index + 2, total, title)

    _closing_slide(presentation, blank, plan, requirements, total)

    _apply_ppt_style(presentation, str(requirements.get("PPT视觉风格", "科技蓝")))
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
