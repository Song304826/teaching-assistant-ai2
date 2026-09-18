from __future__ import annotations

import hashlib
import json
import math
import re
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pptx import Presentation


SUPPORTED_SUFFIXES = {".docx", ".pdf", ".pptx", ".txt", ".md"}
INDEX_FILENAME = ".knowledge_index.json"


@dataclass(frozen=True)
class KnowledgeChunk:
    filename: str
    title: str
    category: str
    source: str
    text: str


@dataclass(frozen=True)
class SearchResult:
    filename: str
    title: str
    category: str
    source: str
    text: str
    score: float


@dataclass(frozen=True)
class SyncReport:
    added: int
    updated: int
    unchanged: int
    removed: int
    failed: int
    documents: int
    chunks: int
    synced_at: str


def _tokens(text: str) -> list[str]:
    normalized = re.sub(r"\s+", "", text.lower())
    chinese = re.findall(r"[\u4e00-\u9fff]+", normalized)
    words = re.findall(r"[a-z0-9]+", normalized)
    result: list[str] = list(words)
    for sequence in chinese:
        result.extend(sequence)
        result.extend(sequence[index : index + 2] for index in range(len(sequence) - 1))
    return result


def _docx_blocks(source) -> list[str]:
    document = Document(source)
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            line = " ｜ ".join(re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells if cell.text.strip())
            if line:
                blocks.append(line)
    return blocks


def _pdf_blocks(source) -> list[str]:
    reader = PdfReader(source)
    return [text for page in reader.pages if (text := (page.extract_text() or "").strip())]


def _pptx_blocks(source) -> list[str]:
    presentation = Presentation(source)
    blocks: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        if texts:
            blocks.append(f"第{index}页｜" + " ｜ ".join(texts))
    return blocks


def _read_blocks(source, suffix: str) -> list[str]:
    suffix = suffix.lower()
    if suffix == ".docx":
        return _docx_blocks(source)
    if suffix == ".pdf":
        return _pdf_blocks(source)
    if suffix == ".pptx":
        return _pptx_blocks(source)
    if suffix in {".txt", ".md"}:
        if isinstance(source, (str, Path)):
            return [Path(source).read_text(encoding="utf-8", errors="ignore")]
        return [source.read().decode("utf-8", errors="ignore")]
    raise ValueError(f"暂不支持 {suffix} 格式")


def _split_blocks(blocks: list[str], max_chars: int = 650) -> list[str]:
    chunks: list[str] = []
    current: list[str] = []
    length = 0
    for raw in blocks:
        block = re.sub(r"\n{3,}", "\n\n", raw).strip()
        if not block:
            continue
        if current and length + len(block) > max_chars:
            chunks.append("\n".join(current))
            current, length = [], 0
        current.append(block)
        length += len(block)
    if current:
        chunks.append("\n".join(current))
    return chunks


def parse_uploaded_document(filename: str, data: bytes, category: str = "教师上传") -> list[KnowledgeChunk]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("仅支持 PDF、DOCX、PPTX、TXT 和 Markdown 资料")
    blocks = _read_blocks(BytesIO(data), suffix)
    if not blocks or not any(block.strip() for block in blocks):
        raise ValueError("没有提取到可检索文字；扫描版 PDF 请先进行 OCR")
    return [
        KnowledgeChunk(filename, Path(filename).stem, category, f"教师本次上传｜{filename}", text)
        for text in _split_blocks(blocks)
    ]


def _rank(query: str, chunks: list[KnowledgeChunk], top_k: int) -> list[SearchResult]:
    if not query.strip() or not chunks:
        return []
    term_counts = [Counter(_tokens(chunk.text + " " + chunk.title + " " + chunk.category)) for chunk in chunks]
    document_frequency: Counter[str] = Counter()
    for counts in term_counts:
        document_frequency.update(counts.keys())
    total_docs = max(len(chunks), 1)
    idf = {token: math.log((total_docs + 1) / (frequency + 1)) + 1 for token, frequency in document_frequency.items()}

    def vector(counts: Counter[str]) -> dict[str, float]:
        total = sum(counts.values()) or 1
        return {token: (count / total) * idf.get(token, math.log(total_docs + 1) + 1) for token, count in counts.items()}

    def cosine(left: dict[str, float], right: dict[str, float]) -> float:
        shared = left.keys() & right.keys()
        numerator = sum(left[token] * right[token] for token in shared)
        left_norm = math.sqrt(sum(value * value for value in left.values()))
        right_norm = math.sqrt(sum(value * value for value in right.values()))
        return numerator / (left_norm * right_norm) if left_norm and right_norm else 0.0

    query_vector = vector(Counter(_tokens(query)))
    scored: list[SearchResult] = []
    for chunk, counts in zip(chunks, term_counts):
        score = cosine(query_vector, vector(counts))
        if score > 0:
            scored.append(SearchResult(**chunk.__dict__, score=score))
    scored.sort(key=lambda item: item.score, reverse=True)
    return scored[:top_k]


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class LocalKnowledgeBase:
    """可增量更新的本地知识库；新增资料无需改动主程序。"""

    def __init__(self, root: Path):
        self.root = Path(root)
        self.documents_root = self.root / "documents"
        self.documents_root.mkdir(parents=True, exist_ok=True)
        self.index_path = self.root / INDEX_FILENAME
        self.chunks: list[KnowledgeChunk] = []
        self.document_count = 0
        self.categories: dict[str, int] = {}
        self.last_sync_report = self.sync()

    def _load_index(self) -> dict:
        if not self.index_path.exists():
            return {"files": {}}
        try:
            content = json.loads(self.index_path.read_text(encoding="utf-8"))
            return content if isinstance(content, dict) else {"files": {}}
        except Exception:
            return {"files": {}}

    def _metadata(self) -> dict[str, dict]:
        manifest_path = self.root / "manifest.json"
        if not manifest_path.exists():
            return {}
        try:
            data = json.loads(manifest_path.read_text(encoding="utf-8"))
            return {str(item.get("filename", "")): item for item in data.get("documents", [])}
        except Exception:
            return {}

    def sync(self, force: bool = False) -> SyncReport:
        previous = self._load_index().get("files", {})
        metadata = self._metadata()
        current: dict[str, dict] = {}
        added = updated = unchanged = failed = 0
        paths = sorted(path for path in self.documents_root.rglob("*") if path.suffix.lower() in SUPPORTED_SUFFIXES)

        for path in paths:
            relative = path.relative_to(self.documents_root).as_posix()
            try:
                digest = _file_hash(path)
                old = previous.get(relative, {})
                if not force and old.get("sha256") == digest and isinstance(old.get("chunks"), list):
                    current[relative] = old
                    unchanged += 1
                    continue
                blocks = _read_blocks(path, path.suffix)
                text_chunks = _split_blocks(blocks)
                if not text_chunks:
                    raise ValueError("没有提取到文字")
                record = metadata.get(path.name, {})
                category = record.get("category") or (path.parent.name if path.parent != self.documents_root else "未分类")
                chunk_records = [
                    asdict(KnowledgeChunk(
                        filename=relative,
                        title=record.get("title", path.stem),
                        category=category,
                        source=record.get("source", f"团队知识库｜{relative}"),
                        text=text,
                    ))
                    for text in text_chunks
                ]
                current[relative] = {"sha256": digest, "chunks": chunk_records}
                if relative in previous:
                    updated += 1
                else:
                    added += 1
            except Exception as exc:
                failed += 1
                current[relative] = {"sha256": "", "chunks": [], "error": str(exc)[:160]}

        removed = len(set(previous) - set(current))
        synced_at = datetime.now().isoformat(timespec="seconds")
        payload = {"version": 1, "synced_at": synced_at, "files": current}
        try:
            self.index_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        except OSError:
            pass

        self.chunks = []
        self.categories = {}
        for item in current.values():
            for raw in item.get("chunks", []):
                try:
                    chunk = KnowledgeChunk(**raw)
                except TypeError:
                    continue
                self.chunks.append(chunk)
                self.categories[chunk.category] = self.categories.get(chunk.category, 0) + 1
        self.document_count = sum(1 for item in current.values() if item.get("chunks"))
        report = SyncReport(added, updated, unchanged, removed, failed, self.document_count, len(self.chunks), synced_at)
        self.last_sync_report = report
        return report

    def search(
        self,
        query: str,
        top_k: int = 5,
        extra_chunks: list[KnowledgeChunk] | None = None,
        categories: list[str] | None = None,
    ) -> list[SearchResult]:
        chunks = self.chunks + list(extra_chunks or [])
        if categories:
            selected = set(categories)
            chunks = [chunk for chunk in chunks if chunk.category in selected or chunk.category in {"综合资料", "未分类"}]
        return _rank(query, chunks, top_k)
